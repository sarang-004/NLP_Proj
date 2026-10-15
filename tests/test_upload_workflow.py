from io import BytesIO
from docx import Document

import numpy as np
import pytest

import app as dashboard_app
from screening.upload_workflow import (
    ResumeUploadError,
    evaluate_uploads,
    extract_uploaded_text,
    extract_jd_minimum_experience,
    make_upload_job,
    structure_uploaded_resume,
)


class FixedEmbeddingModel:
    def encode(self, texts, **kwargs):
        return np.tile(np.array([[1.0, 0.0]]), (len(texts), 1))


def job_fields(**overrides):
    fields = {
        "job_description": "Software engineer. Required skills: Python and SQL.",
        "job_title": "Software engineer",
        "required_skills": "Python, SQL",
        "minimum_experience": "2",
        "preferred_degree": "B.Tech",
        "accepted_degrees": "B.Tech|B.E.",
        "preferred_course": "Computer Science",
        "accepted_courses": "Computer Science|Information Technology",
    }
    fields.update(overrides)
    return fields


def one_page_pdf(text):
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(payload))
        payload.extend(f"{index} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")
    xref = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode("ascii"))
    return bytes(payload)


def test_txt_upload_extraction_and_unsupported_file():
    assert extract_uploaded_text("one.txt", BytesIO(b"Resume text")) == "Resume text"
    with pytest.raises(ResumeUploadError, match="Unsupported"):
        extract_uploaded_text("one.exe", BytesIO(b"data"))
    with pytest.raises(ResumeUploadError, match="empty"):
        extract_uploaded_text("empty.txt", BytesIO(b""))


def test_pdf_docx_and_txt_extract_readable_resume_text():
    pdf_text = extract_uploaded_text("one.pdf", BytesIO(one_page_pdf("Python resume evidence")))
    assert "Python resume evidence" in pdf_text
    document = Document()
    document.add_paragraph("Java resume evidence")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "SQL table evidence"
    docx_content = BytesIO()
    document.save(docx_content)
    docx_text = extract_uploaded_text("one.docx", BytesIO(docx_content.getvalue()))
    assert "Java resume evidence" in docx_text
    assert "SQL table evidence" in docx_text
    with pytest.raises(ResumeUploadError, match="No readable text"):
        extract_uploaded_text("blank.txt", BytesIO(b" \r\n\t"))
    with pytest.raises(ResumeUploadError, match="No readable text"):
        extract_uploaded_text("blank.pdf", BytesIO(one_page_pdf("")))


def test_api_runs_pdf_docx_and_txt_resume_uploads_end_to_end(monkeypatch):
    monkeypatch.setattr(dashboard_app, "_load_upload_model", lambda: FixedEmbeddingModel())
    document = Document()
    document.add_paragraph("Skills: Python")
    docx_content = BytesIO()
    document.save(docx_content)
    samples = [
        ("resume.pdf", one_page_pdf("Skills: Python")),
        ("resume.docx", docx_content.getvalue()),
        ("resume.txt", b"Skills: Python"),
    ]
    for filename, content in samples:
        response = dashboard_app.app.test_client().post("/api/evaluate", data={
            "job_title": "Engineer", "job_description": "Required skills: Python.",
            "resumes": (BytesIO(content), filename),
        }, content_type="multipart/form-data")
        assert response.status_code == 200
        payload = response.get_json()
        assert payload["candidate_count"] == 1
        assert payload["scoring_methodology_version"] == "evidence-v4"
        assert payload["candidates"][0]["skill_evidence"][0]["evidence_status"] == "listed"


@pytest.mark.parametrize(("filename", "content", "error"), [
    ("broken.pdf", b"not a PDF", "Could not read PDF file"),
    ("broken.docx", b"not a DOCX", "Could not read DOCX file"),
])
def test_malformed_document_uploads_fail_without_extracted_text(filename, content, error):
    with pytest.raises(ResumeUploadError, match=error):
        extract_uploaded_text(filename, BytesIO(content))


def test_upload_job_detects_skills_and_experience_from_jd_without_overrides():
    job = make_upload_job({
        "job_title": "AWS Engineer",
        "job_description": "AWS Engineer. Requires 3+ years of experience with AWS, Python, and SQL.",
    })
    assert job["required_skills_list"] == ["AWS", "Python", "SQL"]
    assert job["minimum_total_experience_years"] == 3
    assert job["required_skills_source"] == "job_description"
    assert job["experience_requirement_source"] == "job_description"


def test_experience_parser_uses_lower_bound_for_stated_range():
    assert extract_jd_minimum_experience("At least 2-4 years of relevant experience") == 2
    assert extract_jd_minimum_experience("Experience: 5 years") == 5
    assert extract_jd_minimum_experience("No minimum experience stated") is None
    assert extract_jd_minimum_experience("4+ years developing in Ruby and Java") == 4


def test_upload_job_manual_requirements_override_jd_detection():
    job = make_upload_job(job_fields(
        job_description="AWS Engineer. Requires 5 years experience with AWS.",
        required_skills="Python",
        minimum_experience="2",
    ))
    assert job["required_skills_list"] == ["Python"]
    assert job["minimum_total_experience_years"] == 2
    assert job["required_skills_source"] == "user_override"
    assert job["experience_requirement_source"] == "user_override"


@pytest.mark.parametrize("value", ["NaN", "inf", "-inf"])
def test_upload_job_rejects_nonfinite_experience_override(value):
    with pytest.raises(ValueError, match="finite number"):
        make_upload_job(job_fields(minimum_experience=value))


def test_text_adapter_preserves_missing_demographic_and_experience_values():
    row = structure_uploaded_resume("Skills: Python\nNo dates or age listed.", "Candidate 1")
    assert row["age"] == "NA"
    assert row["total_experience_years"] == "NA"
    assert row["college"] == "NA"


def test_resume_name_skips_summary_heading_and_filename_repairs_split_name():
    from screening.upload_workflow import extract_resume_name

    assert extract_resume_name("SUMMARY\nMIDHUNLAL P\nUdinur, Kasargod") == "MIDHUNLAL P"
    jd = make_upload_job({"job_title": "Engineer", "job_description": "Required skills: Python."})
    candidates = evaluate_uploads(
        [("one", "SUMMARY\nGAUTHAM PRA VEEN\nSkills: Python", "RESUME-GAUTHAM PRAVEEN")],
        jd,
        FixedEmbeddingModel(),
    )
    assert candidates[0]["candidate_name"] == "GAUTHAM PRAVEEN"
    assert candidates[0]["candidate_name_source"] == "filename_confirmed"


def test_resume_adapter_reads_spaced_skills_heading_and_profile_experience():
    row = structure_uploaded_resume(
        "S K I L L S\nManual Testing\nAutomation Testing\nT O O L S / T E C H N O L O G I E S\nJava\n"
        "PROFILE\nI have 7+ years of experience across testing and QA.",
        "Candidate 1",
    )
    skills = set(row["skills_extracted"].split("|"))
    assert "Java" in skills
    assert "Automation" in skills or "Automation Testing" in skills
    assert float(row["total_experience_years"]) == 7


def test_multiple_resumes_share_job_and_are_ranked_with_real_components():
    job = make_upload_job(job_fields())
    uploads = [
        ("Candidate 1", "Skills: Python, SQL\nTotal experience: 5 years\nB.Tech in Computer Science\nProject: Hiring tool"),
        ("Candidate 2", "Skills: Python\nTotal experience: 1 year\nB.E. in Information Technology"),
    ]
    candidates = evaluate_uploads(uploads, job, FixedEmbeddingModel())
    assert [candidate["rank"] for candidate in candidates] == [1, 2]
    assert all(candidate["component_scores"]["skills"] > 0 for candidate in candidates)
    assert all("education_course" in candidate["weighted_contributions"] for candidate in candidates)
    assert all(candidate["degree_course_evidence"]["degree"] != "" for candidate in candidates)
    for candidate in candidates:
        assert candidate["final_score"] == pytest.approx(sum(candidate["weighted_contributions"].values()))
        assert sum(candidate["effective_weights"].values()) == pytest.approx(1.0)
        assert candidate["configured_weights"] == pytest.approx({
            "skills": 0.30, "experience": 0.20, "education_course": 0.15, "projects": 0.15, "semantic": 0.20
        })
        assert candidate["score_denominator_weight"] == pytest.approx(sum(
            candidate["configured_weights"][name] for name, value in candidate["component_scores"].items() if value is not None
        ))
        for name, component_score in candidate["component_scores"].items():
            if component_score is None:
                assert name in candidate["excluded_factors"]
                assert candidate["weighted_contributions"][name] == 0
            else:
                assert candidate["effective_weights"][name] == pytest.approx(
                    candidate["configured_weights"][name] / candidate["score_denominator_weight"]
                    , abs=1e-6
                )
                assert candidate["weighted_contributions"][name] == pytest.approx(
                    component_score * candidate["effective_weights"][name], abs=1e-6
                )


def test_upload_api_handles_invalid_file_and_scores_multiple_files(monkeypatch):
    monkeypatch.setattr(dashboard_app, "_load_upload_model", lambda: FixedEmbeddingModel())
    client = dashboard_app.app.test_client()
    response = client.post("/api/evaluate", data={
        **job_fields(),
        "resumes": [
            (BytesIO(b"Skills: Python, SQL\nTotal experience: 5 years\nB.Tech in Computer Science"), "one.txt"),
            (BytesIO(b"bad"), "bad.exe"),
            (BytesIO(b"Skills: Python\nTotal experience: 1 year\nB.E. in Information Technology"), "two.txt"),
        ],
    }, content_type="multipart/form-data")
    assert response.status_code == 200
    body = response.get_json()
    assert body["candidate_count"] == 2
    assert len(body["file_errors"]) == 1
    assert body["persisted"] is False


def test_upload_api_returns_jd_derived_criteria(monkeypatch):
    monkeypatch.setattr(dashboard_app, "_load_upload_model", lambda: FixedEmbeddingModel())
    client = dashboard_app.app.test_client()
    response = client.post("/api/evaluate", data={
        "job_title": "AWS Engineer",
        "job_description": "AWS Engineer. Requires at least 3 years experience with AWS and Python.",
        "resumes": (BytesIO(b"Skills: AWS, Python\nTotal experience: 4 years"), "one.txt"),
    }, content_type="multipart/form-data")
    assert response.status_code == 200
    body = response.get_json()
    assert body["criteria"]["required_skills"] == ["AWS", "Python"]
    assert body["criteria"]["minimum_experience_years"] == 3
    assert any("cannot verify how many years" in warning for warning in body["warnings"])


def test_upload_api_accepts_pdf_job_description(monkeypatch):
    monkeypatch.setattr(dashboard_app, "_load_upload_model", lambda: FixedEmbeddingModel())
    real_extract = dashboard_app.extract_uploaded_text

    def extract(filename, stream):
        if filename == "job.pdf":
            return "AWS Engineer. Qualifications: 4+ years developing in Ruby and Java."
        return real_extract(filename, stream)

    monkeypatch.setattr(dashboard_app, "extract_uploaded_text", extract)
    client = dashboard_app.app.test_client()
    response = client.post("/api/evaluate", data={
        "jd_source": "pdf",
        "job_title": "AWS Engineer",
        "job_description_pdf": (BytesIO(b"pdf placeholder"), "job.pdf"),
        "resumes": (BytesIO(b"Skills: Java\nTotal experience: 7 years"), "one.txt"),
    }, content_type="multipart/form-data")
    assert response.status_code == 200
    body = response.get_json()
    assert body["criteria"]["minimum_experience_years"] == 4
    assert "Ruby" in body["criteria"]["required_skills"]


def test_upload_api_reports_missing_files_and_job_fields():
    client = dashboard_app.app.test_client()
    assert client.post("/api/evaluate", data={}).status_code == 400
    response = client.post("/api/evaluate", data={"resumes": (BytesIO(b"x"), "resume.txt")})
    assert response.status_code == 400
    assert "job description" in response.get_json()["error"].casefold()


@pytest.mark.parametrize(("filename", "content", "message"), [
    ("resume.exe", b"not a resume", "Unsupported file type"),
    ("empty.txt", b"", "uploaded file is empty"),
    ("broken.pdf", b"not a PDF", "Could not read PDF file"),
    ("broken.docx", b"not a DOCX", "Could not read DOCX file"),
])
def test_api_rejects_invalid_resume_without_fabricated_candidate_score(monkeypatch, filename, content, message):
    monkeypatch.setattr(dashboard_app, "_load_upload_model", lambda: FixedEmbeddingModel())
    response = dashboard_app.app.test_client().post("/api/evaluate", data={
        "job_title": "Engineer", "job_description": "Required skills: Python.",
        "resumes": (BytesIO(content), filename),
    }, content_type="multipart/form-data")
    assert response.status_code == 400
    payload = response.get_json()
    assert payload["error"] == "No resumes could be evaluated."
    assert message.casefold() in payload["file_errors"][0]["message"].casefold()
    assert "candidates" not in payload


def test_api_rejects_missing_or_invalid_job_description_pdf(monkeypatch):
    client = dashboard_app.app.test_client()
    resume = (BytesIO(b"Skills: Python"), "resume.txt")
    missing = client.post("/api/evaluate", data={
        "jd_source": "pdf", "job_title": "Engineer", "resumes": resume,
    }, content_type="multipart/form-data")
    assert missing.status_code == 400
    assert "Choose a PDF job description" in missing.get_json()["error"]
    wrong_type = client.post("/api/evaluate", data={
        "jd_source": "pdf", "job_title": "Engineer",
        "job_description_pdf": (BytesIO(b"JD"), "jd.txt"),
        "resumes": (BytesIO(b"Skills: Python"), "resume.txt"),
    }, content_type="multipart/form-data")
    assert wrong_type.status_code == 400
    assert "must be a PDF" in wrong_type.get_json()["error"]
    malformed = client.post("/api/evaluate", data={
        "jd_source": "pdf", "job_title": "Engineer",
        "job_description_pdf": (BytesIO(b"not a PDF"), "jd.pdf"),
        "resumes": (BytesIO(b"Skills: Python"), "resume.txt"),
    }, content_type="multipart/form-data")
    assert malformed.status_code == 400
    assert "Could not read the job description PDF" in malformed.get_json()["error"]
